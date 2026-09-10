import os
import json
from pathlib import Path
from typing import Dict, Any, Optional, List
from datetime import datetime
from instrument_interface import InstrumentData, AcquisitionResult

# 尝试导入openpyxl库
try:
    from openpyxl import load_workbook
    from openpyxl.styles import Font, Alignment, Border, Side
    from openpyxl.chart import LineChart, Reference
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


class ReportGenerator:
    def __init__(self, template_dir: str = "templates"):
        self.template_dir = Path(template_dir)
        self.template_dir.mkdir(parents=True, exist_ok=True)
        self._ensure_templates_exist()

    def _ensure_templates_exist(self):
        """确保模板文件存在"""
        # 创建默认的Excel模板
        default_template = self.template_dir / "default_report_template.xlsx"
        if not default_template.exists():
            self._create_default_template(default_template)

    def _create_default_template(self, template_path: Path):
        """创建默认的Excel模板"""
        if not OPENPYXL_AVAILABLE:
            print("openpyxl is not installed, cannot create template")
            return

        # 创建一个简单的Excel工作簿作为模板
        from openpyxl import Workbook
        wb = Workbook()
        
        # 移除默认的工作表
        wb.remove(wb.active)
        
        # 创建摘要工作表
        summary_ws = wb.create_sheet("摘要")
        summary_ws.title = "摘要"
        
        # 设置标题
        summary_ws['A1'] = "仪器数据采集报告"
        summary_ws['A1'].font = Font(bold=True, size=16)
        summary_ws.merge_cells('A1:D1')
        summary_ws['A1'].alignment = Alignment(horizontal='center')
        
        # 设置报告信息
        summary_ws['A3'] = "报告生成时间:"
        summary_ws['B3'] = "{report_time}"
        summary_ws['A4'] = "采集时间范围:"
        summary_ws['B4'] = "{start_time} 至 {end_time}"
        summary_ws['A5'] = "采集仪器数量:"
        summary_ws['B5'] = "{instrument_count}"
        
        # 设置表头
        summary_ws['A7'] = "仪器ID"
        summary_ws['B7'] = "仪器类型"
        summary_ws['C7'] = "采集次数"
        summary_ws['D7'] = "状态"
        
        # 设置表头样式
        for cell in summary_ws['A7:D7']:
            for c in cell:
                c.font = Font(bold=True)
                c.border = Border(left=Side(style='thin'), right=Side(style='thin'), 
                                top=Side(style='thin'), bottom=Side(style='thin'))
                c.alignment = Alignment(horizontal='center')
        
        # 创建数据工作表
        data_ws = wb.create_sheet("数据")
        data_ws.title = "数据"
        
        # 设置数据表格
        data_ws['A1'] = "采集数据"
        data_ws['A1'].font = Font(bold=True, size=14)
        data_ws.merge_cells('A1:F1')
        
        data_ws['A3'] = "时间"
        data_ws['B3'] = "仪器ID"
        data_ws['C3'] = "参数类型"
        data_ws['D3'] = "参数名称"
        data_ws['E3'] = "数值"
        data_ws['F3'] = "单位"
        
        # 设置表头样式
        for cell in data_ws['A3:F3']:
            for c in cell:
                c.font = Font(bold=True)
                c.border = Border(left=Side(style='thin'), right=Side(style='thin'), 
                                top=Side(style='thin'), bottom=Side(style='thin'))
                c.alignment = Alignment(horizontal='center')
        
        # 创建图表工作表
        chart_ws = wb.create_sheet("图表")
        chart_ws.title = "图表"
        chart_ws['A1'] = "数据趋势图"
        chart_ws['A1'].font = Font(bold=True, size=14)
        
        # 保存模板
        wb.save(template_path)
        print(f"Created default template at: {template_path}")

    def generate_report(self, results: List[AcquisitionResult], 
                      output_path: str, 
                      template_name: str = "default_report_template.xlsx") -> bool:
        """生成Excel报告"""
        if not OPENPYXL_AVAILABLE:
            print("openpyxl is not installed, cannot generate report")
            return False

        try:
            # 加载模板
            template_path = self.template_dir / template_name
            if not template_path.exists():
                print(f"Template not found: {template_path}")
                return False
            
            wb = load_workbook(template_path)
            
            # 处理摘要工作表
            if "摘要" in wb.sheetnames:
                summary_ws = wb["摘要"]
                self._fill_summary(summary_ws, results)
            
            # 处理数据工作表
            if "数据" in wb.sheetnames:
                data_ws = wb["数据"]
                self._fill_data(data_ws, results)
            
            # 处理图表工作表
            if "图表" in wb.sheetnames:
                chart_ws = wb["图表"]
                self._fill_chart(chart_ws, results)
            
            # 保存报告
            output_path = Path(output_path)
            output_path.parent.mkdir(parents=True, exist_ok=True)
            wb.save(output_path)
            print(f"Report generated successfully: {output_path}")
            return True
        except Exception as e:
            print(f"Error generating report: {e}")
            return False

    def _fill_summary(self, ws, results: List[AcquisitionResult]):
        """填充摘要工作表"""
        # 生成报告时间
        report_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        ws['B3'] = report_time
        
        # 计算时间范围
        if results:
            timestamps = [r.sync_timestamp for r in results]
            start_time = min(timestamps)
            end_time = max(timestamps)
            start_str = datetime.fromtimestamp(start_time).strftime("%Y-%m-%d %H:%M:%S")
            end_str = datetime.fromtimestamp(end_time).strftime("%Y-%m-%d %H:%M:%S")
            ws['B4'] = f"{start_str} 至 {end_str}"
        
        # 统计仪器数量
        instrument_ids = set()
        for result in results:
            instrument_ids.update(result.results.keys())
        ws['B5'] = len(instrument_ids)
        
        # 填充仪器信息
        row = 8
        for inst_id in instrument_ids:
            # 统计该仪器的采集次数
            count = 0
            for result in results:
                if inst_id in result.results:
                    count += 1
            
            # 获取仪器类型
            inst_type = "未知"
            for result in results:
                if inst_id in result.results:
                    inst_type = result.results[inst_id].instrument_type.value
                    break
            
            # 填充数据
            ws[f'A{row}'] = inst_id
            ws[f'B{row}'] = inst_type
            ws[f'C{row}'] = count
            ws[f'D{row}'] = "正常"
            
            # 设置边框
            for col in ['A', 'B', 'C', 'D']:
                cell = ws[f'{col}{row}']
                cell.border = Border(left=Side(style='thin'), right=Side(style='thin'), 
                                  top=Side(style='thin'), bottom=Side(style='thin'))
            
            row += 1

    def _fill_data(self, ws, results: List[AcquisitionResult]):
        """填充数据工作表"""
        row = 4
        
        for result in results:
            timestamp = result.sync_timestamp
            datetime_str = result.sync_datetime
            
            for instrument_id, data in result.results.items():
                # 处理功率测量数据
                for measurement in data.power_measurements:
                    ws[f'A{row}'] = datetime_str
                    ws[f'B{row}'] = instrument_id
                    ws[f'C{row}'] = "功率测量"
                    ws[f'D{row}'] = f"电压 (通道{measurement.element_id})"
                    ws[f'E{row}'] = measurement.voltage
                    ws[f'F{row}'] = "V"
                    row += 1
                    
                    ws[f'A{row}'] = datetime_str
                    ws[f'B{row}'] = instrument_id
                    ws[f'C{row}'] = "功率测量"
                    ws[f'D{row}'] = f"电流 (通道{measurement.element_id})"
                    ws[f'E{row}'] = measurement.current
                    ws[f'F{row}'] = "A"
                    row += 1
                    
                    ws[f'A{row}'] = datetime_str
                    ws[f'B{row}'] = instrument_id
                    ws[f'C{row}'] = "功率测量"
                    ws[f'D{row}'] = f"有功功率 (通道{measurement.element_id})"
                    ws[f'E{row}'] = measurement.active_power
                    ws[f'F{row}'] = "W"
                    row += 1
                
                # 处理波形数据（仅记录通道信息）
                for channel in data.waveform_channels:
                    ws[f'A{row}'] = datetime_str
                    ws[f'B{row}'] = instrument_id
                    ws[f'C{row}'] = "波形数据"
                    ws[f'D{row}'] = f"通道 {channel.channel_name}"
                    ws[f'E{row}'] = f"{len(channel.voltage_data)} 数据点"
                    ws[f'F{row}'] = channel.unit
                    row += 1

    def _fill_chart(self, ws, results: List[AcquisitionResult]):
        """填充图表工作表"""
        # 简单的图表生成逻辑
        # 这里只是示例，实际项目中可能需要更复杂的图表生成
        ws['A3'] = "注：图表功能需要根据实际数据类型和需求进行定制"
        ws['A4'] = "当前版本仅提供数据表格，图表功能将在后续版本中完善"

    def generate_single_instrument_report(self, data: InstrumentData, 
                                         output_path: str, 
                                         template_name: str = "default_report_template.xlsx") -> bool:
        """生成单台仪器的报告"""
        # 创建一个包含单个结果的列表
        from instrument_interface import AcquisitionResult
        result = AcquisitionResult(
            sequence_number=1,
            sync_timestamp=data.timestamp,
            sync_datetime=data.datetime_str,
            results={data.instrument_id: data},
            errors={},
            acquisition_time_ms=0
        )
        return self.generate_report([result], output_path, template_name)
